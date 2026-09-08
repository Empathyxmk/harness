import os
import sys
import pytest
import openpyxl
import json

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '../src')))
from convert_excel_to_json import convert_excel_to_json

THIS_DIR = os.path.dirname(os.path.abspath(__file__))
TEST_XLSX = os.path.abspath(os.path.join(THIS_DIR, '../tests/test-data-2.xlsx'))
TEST_XLSX2 = os.path.abspath(os.path.join(THIS_DIR, '../tests/test-data.xlsx'))

def test_load_module_and_is_function():
    assert callable(convert_excel_to_json)

def test_throw_if_no_source_file():
    with pytest.raises(Exception):
        convert_excel_to_json({})

def test_parse_basic_sheet_to_json():
    result = convert_excel_to_json({'sourceFile': TEST_XLSX})
    assert isinstance(result, dict)
    ws_names = list(result)
    assert len(ws_names) > 0
    assert isinstance(result[ws_names[0]][0], dict)

def test_parse_only_specified_sheet():
    wb = openpyxl.load_workbook(TEST_XLSX)
    sheet_name = wb.sheetnames[-1]
    result = convert_excel_to_json({'sourceFile': TEST_XLSX, 'sheets': [sheet_name]})
    assert set(result) == set([sheet_name])

def test_return_empty_for_nonexistent_sheet():
    result = convert_excel_to_json({'sourceFile': TEST_XLSX, 'sheets': [{'name':'Sheet_DOES_NOT_EXIST'}]})
    assert result['Sheet_DOES_NOT_EXIST'] == []

def test_respect_include_empty_lines_false():
    res = convert_excel_to_json({'sourceFile': TEST_XLSX})
    ws = list(res)[0]
    assert all(row is not None for row in res[ws])

def test_include_empty_rows_true():
    res = convert_excel_to_json({'sourceFile': TEST_XLSX, 'includeEmptyLines': True})
    ws = list(res)[0]
    assert isinstance(res[ws], list)

def test_range_and_column_to_key():
    res = convert_excel_to_json({'sourceFile': TEST_XLSX, 'range':'B2:C4','columnToKey':{'B':'bCol','C':'cCol'}})
    ws = list(res)[0]
    for row in res[ws]:
        assert set(row.keys()).issubset({'bCol','cCol'})

def test_append_data():
    wb = openpyxl.load_workbook(TEST_XLSX2)
    sheet_name = wb.sheetnames[-1]
    res = convert_excel_to_json({'sourceFile': TEST_XLSX2, 'sheets': [{'name':sheet_name, 'appendData': {'bar':321}}]})
    assert res[sheet_name][0]['bar'] == 321

def test_support_stubs():
    res = convert_excel_to_json({'sourceFile': TEST_XLSX, 'sheetStubs': True})
    ws = list(res)[0]
    assert isinstance(res[ws], list)

def test_handle_config_as_json_string():
    config_str = json.dumps({'sourceFile': TEST_XLSX})
    res = convert_excel_to_json(config_str)
    ws = list(res)[0]
    assert isinstance(res[ws], list)

def test_column_to_key_wildcard():
    res = convert_excel_to_json({'sourceFile': TEST_XLSX, 'columnToKey': {'*':'publiccol'}})
    ws = list(res)[0]
    assert isinstance(res[ws][0], dict)

def test_throw_with_invalid_json_string():
    with pytest.raises(Exception):
        convert_excel_to_json('{not json}')