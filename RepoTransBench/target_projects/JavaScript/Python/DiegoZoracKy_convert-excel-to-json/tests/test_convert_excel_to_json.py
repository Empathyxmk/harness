import os
import sys
import pytest
import openpyxl
import json

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '../src')))
from convert_excel_to_json import convert_excel_to_json

THIS_DIR = os.path.dirname(os.path.abspath(__file__))
TEST_XLSX = os.path.join(THIS_DIR, "test-data.xlsx")
TEST_XLSX2 = os.path.join(THIS_DIR, "test-data-2.xlsx")

def test_loads_without_error():
    assert callable(convert_excel_to_json)

def test_throws_without_source_file():
    with pytest.raises(Exception):
        convert_excel_to_json({})

def test_parse_basic_sheet_to_json():
    result = convert_excel_to_json({'sourceFile': TEST_XLSX})
    assert isinstance(result, dict)
    ws_names = list(result)
    assert len(ws_names) > 0
    assert isinstance(result[ws_names[0]][0], dict)

def test_parse_only_specified_sheet():
    wb = openpyxl.load_workbook(TEST_XLSX)
    sheet_name = wb.sheetnames[0]
    result = convert_excel_to_json({'sourceFile': TEST_XLSX, 'sheets': [sheet_name]})
    assert list(result.keys()) == [sheet_name]

def test_return_empty_for_nonexistent_sheet():
    result = convert_excel_to_json({'sourceFile': TEST_XLSX, 'sheets': [{'name': 'NonExistent'}]})
    assert result['NonExistent'] == []

def test_include_empty_lines_false_default():
    res = convert_excel_to_json({'sourceFile': TEST_XLSX})
    ws = list(res.keys())[0]
    assert all(row is not None for row in res[ws])

def test_include_empty_lines_true_includes_blanks():
    res = convert_excel_to_json({'sourceFile': TEST_XLSX, 'includeEmptyLines': True})
    ws = list(res.keys())[0]
    assert isinstance(res[ws], list)

def test_range_option_and_column_to_key():
    res = convert_excel_to_json({'sourceFile': TEST_XLSX, 'range': 'A1:B3', 'columnToKey': {'A': 'colA', 'B': 'colB'}})
    ws = list(res.keys())[0]
    for row in res[ws]:
        assert all(k in ['colA', 'colB'] for k in row.keys())

def test_support_append_data_property():
    wb = openpyxl.load_workbook(TEST_XLSX2)
    sheet_name = wb.sheetnames[0]
    res = convert_excel_to_json({'sourceFile': TEST_XLSX2, 'sheets': [{'name': sheet_name, 'appendData': {'foo': 123}}]})
    assert res[sheet_name][0].get('foo') == 123

def test_support_sheet_stubs():
    res = convert_excel_to_json({'sourceFile': TEST_XLSX, 'sheetStubs': True})
    ws = list(res.keys())[0]
    assert isinstance(res[ws], list)

def test_config_as_json_string():
    config_str = json.dumps({'sourceFile': TEST_XLSX})
    res = convert_excel_to_json(config_str)
    ws = list(res.keys())[0]
    assert isinstance(res[ws], list)

def test_column_to_key_wildcard():
    res = convert_excel_to_json({'sourceFile': TEST_XLSX, 'columnToKey': {'*': 'anycol'}})
    ws = list(res.keys())[0]
    assert isinstance(res[ws][0], dict)

def test_throw_invalid_json_config_string():
    with pytest.raises(Exception):
        convert_excel_to_json('{invalid}')